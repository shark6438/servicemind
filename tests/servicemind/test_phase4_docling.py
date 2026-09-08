"""Slice: docling/PDF/Office ingest routing to the suffix-aware parser.

Offline assertions (no layout-model snapshot, no network):

1. ``AttachmentSource`` turns a real PDF file into a ``KnowledgeDocument`` whose
   content is extractable plain text and whose content-digest invariant holds.
2. ``StructureParser.parse_file`` routes by suffix and degrades to the pypdf
   fallback when the docling PDF model is unreachable (the predictable offline box),
   and to python-docx for .docx -- never fabricating structure.
3. ``EnterpriseRAG.ingest`` honours ``metadata["source_file"]`` and reaches
   ``parse_file`` instead of the content-prefix sniff, so real binaries index.

The docling primary path is exercised live by scripts/verify_phase4_docling.py against
the real PagerDuty training PDF; here every docling attempt is forced to fail so the
fallback contract is asserted deterministically offline.
"""

import shutil
from pathlib import Path
from uuid import UUID

import pytest

from servicemind.domain.knowledge import (
    AuthorityLevel,
    CorpusScope,
    KnowledgeACL,
    KnowledgeDocument,
    KnowledgeProvenance,
)
from servicemind.rag.models import DeterministicEmbeddingProvider
from servicemind.rag.parsing import BlockType, structure_parser
from servicemind.rag.service import EnterpriseRAG
from servicemind.rag.sources import AttachmentSource

TENANT = UUID("11111111-1111-4111-8111-111111111111")
#: Tracked real PDF (project-owned sample) used as the hermetic binary fixture.
ACME_PDF = Path("data/AcmeTech_Employee_Handbook.pdf")


def _raise_docling(path, document):  # pragma: no cover - body is the point
    raise RuntimeError("offline: docling layout-model snapshot not cached")


def _document(content: str, *, record_id: str, metadata: dict | None = None) -> KnowledgeDocument:
    return KnowledgeDocument(
        title=record_id,
        content=content,
        document_type="attached_document",
        language="en",
        metadata=metadata or {},
        acl=KnowledgeACL(corpus_scope=CorpusScope.TENANT, tenant_id=TENANT),
        provenance=KnowledgeProvenance(
            source="file_attachment",
            source_version="v1",
            source_uri=f"file://{record_id}",
            source_record_id=record_id,
            license="tenant-owned",
            authority_level=AuthorityLevel.INTERNAL_KNOWLEDGE,
            content_hash=KnowledgeDocument.content_digest(content),
        ),
    )


def test_attachment_source_loads_real_pdf_with_valid_invariant(tmp_path) -> None:
    shutil.copy(ACME_PDF, tmp_path / "handbook.pdf")
    import asyncio

    docs = asyncio.run(AttachmentSource(tmp_path, tenant_id=TENANT).load())
    assert len(docs) == 1
    doc = docs[0]
    assert doc.document_type == "attached_document"
    assert doc.metadata["file_format"] == ".pdf"
    assert doc.metadata["source_file"].endswith("handbook.pdf")
    assert "AcmeTech" in doc.content
    assert doc.provenance.content_hash == KnowledgeDocument.content_digest(doc.content)
    assert doc.acl.corpus_scope is CorpusScope.TENANT and doc.acl.tenant_id == TENANT


def test_parse_file_pdf_falls_back_to_pypdf_when_docling_unavailable(monkeypatch) -> None:
    content = structure_parser.file_text(ACME_PDF)
    doc = _document(content, record_id="handbook.pdf", metadata={"source_file": str(ACME_PDF)})
    monkeypatch.setattr(structure_parser, "_docling_blocks", _raise_docling)
    blocks = structure_parser.parse_file(ACME_PDF, doc)
    assert blocks
    joined = "\n".join(block.content for block in blocks)
    assert "AcmeTech" in joined  # pypdf text made it to real blocks
    assert all(block.metadata.get("parser") == "pypdf" for block in blocks)
    assert all(block.block_type is BlockType.PARAGRAPH for block in blocks)


def test_parse_file_docx_falls_back_to_python_docx(monkeypatch, tmp_path) -> None:
    from docx import Document

    target = tmp_path / "note.docx"
    docx = Document()
    docx.add_heading("Maintenance Window Policy", level=1)
    docx.add_paragraph("Scheduled maintenance must be announced to all tenants 48 hours ahead.")
    docx.save(target)
    content = structure_parser.file_text(target)
    assert "Maintenance Window Policy" in content
    doc = _document(content, record_id="note.docx", metadata={"source_file": str(target)})
    monkeypatch.setattr(structure_parser, "_docling_blocks", _raise_docling)
    blocks = structure_parser.parse_file(target, doc)
    kinds = [block.block_type for block in blocks]
    assert BlockType.HEADING in kinds
    assert BlockType.PARAGRAPH in kinds
    assert "48 hours" in " ".join(block.content for block in blocks)


def test_parse_file_markdown_still_parses_content_by_suffix(tmp_path) -> None:
    md = tmp_path / "guide.md"
    md.write_text("# VPN triage\n\nCheck the gateway first.", encoding="utf-8")
    doc = _document(md.read_text(encoding="utf-8"), record_id="guide.md")
    blocks = structure_parser.parse_file(md, doc)
    assert blocks[0].block_type is BlockType.HEADING
    assert blocks[0].content == "VPN triage"


class _RecordingIndex:
    def __init__(self) -> None:
        self.parents_seen: list[list] = []
        self.children_seen: list[list] = []

    async def replace_document(self, tenant_id, document, parents, children, embedding) -> None:
        assert tenant_id == TENANT
        self.parents_seen.append(parents)
        self.children_seen.append(children)

    async def refresh(self, tenant_id) -> None:
        assert tenant_id == TENANT

    async def publish(self, tenant_id, embedding) -> None:
        assert tenant_id == TENANT


class _NoopRepository:
    async def is_current(self, tenant_id, document) -> bool:
        return False

    async def replace(self, tenant_id, document, parents, children):
        return document, parents, children

    async def mark_indexed(self, tenant_id, document_id) -> None:
        return None

    async def count_pending(self, tenant_id) -> int:
        return 0


@pytest.mark.asyncio
async def test_ingest_routes_source_file_to_parse_file(monkeypatch, tmp_path) -> None:
    """A PDF document tagged with source_file must NOT be misparsed as markdown."""
    shutil.copy(ACME_PDF, tmp_path / "handbook.pdf")
    docs = await AttachmentSource(tmp_path, tenant_id=TENANT).load()
    monkeypatch.setattr(structure_parser, "_docling_blocks", _raise_docling)
    # Guard: if ingest ever sniffs content as markdown, the PDF text would be treated
    # as markdown headings/tables -- fail loudly so the routing regression surfaces.
    markdown_calls = []

    original = structure_parser.parse_markdown
    monkeypatch.setattr(
        structure_parser,
        "parse_markdown",
        lambda document: markdown_calls.append(True) or original(document),
    )
    index = _RecordingIndex()
    rag = EnterpriseRAG(
        index=index,  # type: ignore[arg-type]
        embedding=DeterministicEmbeddingProvider(),
        reranker=DeterministicEmbeddingProvider(),  # unused by ingest
        repository=_NoopRepository(),  # type: ignore[arg-type]
    )
    totals = await rag.ingest(TENANT, docs)
    assert totals["documents"] == 1
    assert index.parents_seen and index.parents_seen[0]
    assert index.children_seen and index.children_seen[0]
    # parse_file (docling-with-fallback), not the content-prefix sniff, drove the parse.
    assert markdown_calls == []
