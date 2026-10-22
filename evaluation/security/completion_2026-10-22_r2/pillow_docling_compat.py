import json
import tempfile
from io import BytesIO
from pathlib import Path

import PIL
from docling.document_converter import DocumentConverter
from docx import Document
from PIL import Image

with tempfile.TemporaryDirectory(prefix="servicemind-pillow-compat-") as directory:
    root = Path(directory)
    image = Image.new("RGB", (64, 64), (0, 128, 255))
    for suffix in ("png", "jpg", "webp", "tiff"):
        target = root / f"allowed.{suffix}"
        image.save(target)
        with Image.open(target) as loaded:
            loaded.load()
            assert loaded.size == (64, 64)
    buffer = BytesIO()
    image.save(buffer, format="PNG")
    buffer.seek(0)
    document = Document()
    document.add_heading("Compatibility fixture", 0)
    document.add_paragraph("LOCAL_DOCLING_PILLOW_OK")
    document.add_picture(buffer)
    target = root / "fixture.docx"
    document.save(target)
    converted = DocumentConverter().convert(target).document.export_to_markdown()
    assert "LOCAL_DOCLING_PILLOW_OK" in converted.replace("\\_", "_")
    result = {
        "pillow": PIL.__version__,
        "formats": ["PNG", "JPEG", "WEBP", "TIFF"],
        "docling_docx_with_image": "PASS",
        "verdict": "PASS",
    }
    print(json.dumps(result))
